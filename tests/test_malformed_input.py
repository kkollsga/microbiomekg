"""A raw file that is present but not the shape the prep reads is a *skip
with a reason*, the same channel as an absent file — never an exit that ends
the whole build.

The preps that read published workbooks find their header rows by cell text
and check derived rules against the paper's own numbers. Before this test,
every one of those checks was a ``SystemExit`` raised from inside ``run()``;
the build catches only ``MissingInput``, so one renamed header in one
supplementary table ended a multi-minute build with a bare message, no
report and no census. Each case here corrupts a copy of a real mini fixture
and expects :class:`~microbiomekg.rawdata.MalformedInput`, which the build
reports as "present but malformed" beside the absent ones.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest
from conftest import TAXDUMP_MINI

from microbiomekg import pipeline
from microbiomekg.preps import prep_maier2018, prep_njc19, prep_zimmermann2019
from microbiomekg.rawdata import MalformedInput, MissingInput
from microbiomekg.tables import Frames

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _rename_cell(path: Path, sheet: str, old: str, new: str) -> None:
    """Rename the first cell in ``sheet`` whose stripped text is ``old``."""
    book = openpyxl.load_workbook(path)
    ws = book[sheet]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == old:
                cell.value = new
                book.save(path)
                return
    raise AssertionError(f"{path.name}/{sheet}: no cell {old!r} to rename")


def _rename_sheet(path: Path, old: str, new: str) -> None:
    book = openpyxl.load_workbook(path)
    book[old].title = new
    book.save(path)


def test_njc19_with_its_header_row_renamed_is_malformed_not_fatal(tmp_path):
    xlsx = tmp_path / prep_njc19.RAW_INPUTS[0].rsplit("/", 1)[1]
    shutil.copy(FIXTURES / "njc19_mini" / xlsx.name, xlsx)
    _rename_cell(xlsx, prep_njc19.SHEET, prep_njc19.HEADER_CELL, "Activity")
    with pytest.raises(MalformedInput) as bad:
        prep_njc19.run(tmp_path, Frames(), xlsx=xlsx, taxdump=TAXDUMP_MINI)
    assert prep_njc19.HEADER_CELL in str(bad.value)
    assert isinstance(bad.value, MissingInput)  # the build's one skip channel


def test_njc19_with_its_sheet_renamed_names_the_sheet(tmp_path):
    xlsx = tmp_path / "njc19.xlsx"
    shutil.copy(FIXTURES / "njc19_mini" / "41597_2020_516_MOESM1_ESM.xlsx", xlsx)
    _rename_sheet(xlsx, prep_njc19.SHEET, "Table 5")
    with pytest.raises(MalformedInput) as bad:
        prep_njc19.run(tmp_path, Frames(), xlsx=xlsx, taxdump=TAXDUMP_MINI)
    assert prep_njc19.SHEET in str(bad.value) and "Table 5" in str(bad.value)


def test_maier2018_with_the_species_header_renamed_is_malformed_not_fatal(tmp_path):
    tables = tmp_path / "maier2018"
    shutil.copytree(FIXTURES / "maier2018_mini", tables)
    name, sheet = prep_maier2018.WORKBOOKS[2]
    _rename_cell(tables / name, sheet, prep_maier2018.SPECIES_HEADER_CELL, "NT db")
    with pytest.raises(MalformedInput) as bad:
        prep_maier2018.run(tmp_path, Frames(), tables=tables, taxdump=TAXDUMP_MINI)
    assert prep_maier2018.SPECIES_HEADER_CELL in str(bad.value)


def test_zimmermann2019_with_the_strain_header_renamed_is_malformed_not_fatal(
    tmp_path,
):
    tables = tmp_path / "zimmermann2019"
    shutil.copytree(FIXTURES / "zimmermann2019_mini", tables)
    _rename_cell(
        tables / prep_zimmermann2019.WORKBOOK,
        prep_zimmermann2019.SHEETS["strains"],
        prep_zimmermann2019.STRAIN_HEADER_CELL,
        "Strain",
    )
    with pytest.raises(MalformedInput) as bad:
        prep_zimmermann2019.run(tmp_path, Frames(), tables=tables, taxdump=TAXDUMP_MINI)
    assert prep_zimmermann2019.STRAIN_HEADER_CELL in str(bad.value)


class _Stub:
    def __init__(self, exc: BaseException | None):
        self.exc = exc

    def run(self, raw, store, **_):
        if self.exc is not None:
            raise self.exc
        return {}


def test_the_build_records_a_malformed_source_as_a_skip_with_its_reason(
    tmp_path, monkeypatch, capsys
):
    """``run_prep`` returns the reason for a skip and ``None`` for a run;
    ``prepare`` keeps the reason per source, so the census can say *why* a
    source is not in the graph and not only that it is not."""
    stubs = {
        "maier2018": _Stub(MalformedInput("table 2: no header row 'NT data base'")),
        "njc19": _Stub(MissingInput("no workbook at raw/njc19")),
        "card": _Stub(None),
    }
    monkeypatch.setattr(pipeline, "prep_module", lambda name: stubs[name])
    store = Frames()
    gates = frozenset()
    assert (
        pipeline.run_prep("card", tmp_path, store, scope="microbial", gates=gates)
        is None
    )
    reason = pipeline.run_prep(
        "maier2018", tmp_path, store, scope="microbial", gates=gates
    )
    assert reason == "table 2: no header row 'NT data base'"
    assert "present but malformed" in capsys.readouterr().out
    reason = pipeline.run_prep("njc19", tmp_path, store, scope="microbial", gates=gates)
    assert reason == "no workbook at raw/njc19"
    assert "absent or not opted into" in capsys.readouterr().out


def test_prepare_returns_the_skip_reasons_by_source(tmp_path, monkeypatch):
    bad = MalformedInput("the sheet is not the paper's")
    monkeypatch.setattr(pipeline, "prep_module", lambda name: _Stub(bad))
    store, loaded, skipped = pipeline.prepare(tmp_path)
    assert loaded == [] and "taxon" not in store
    every = {
        p.stem.removeprefix("prep_") for p in pipeline.order_preps(pipeline.PREPS_DIR)
    }
    assert set(skipped) == every
    assert all(reason == str(bad) for reason in skipped.values())


def test_an_unexpected_exception_from_a_prep_still_propagates(tmp_path, monkeypatch):
    """Only the two input classes are skips. A defect in a prep is a defect."""
    monkeypatch.setattr(pipeline, "prep_module", lambda name: _Stub(ValueError("bug")))
    with pytest.raises(ValueError):
        pipeline.run_prep(
            "card", tmp_path, Frames(), scope="microbial", gates=frozenset()
        )
