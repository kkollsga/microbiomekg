#!/usr/bin/env python3
"""Zimmermann 2019's 271 x 76 metabolism screen -> the taxon->drug edges D8 was
still missing.

``data/raw/drug_screens/zimmermann2019/41586_2019_1291_MOESM1_ESM.xlsx`` is the
single 40 MB supplementary workbook of Zimmermann et al., *Nature* 570:462-467
(2019), PMID 31158845, from the Springer static host. It holds 21 sheets; four
are read here:

* ``Supplementary Table 3`` — **the screen**: 271 drugs down, 80 measured
  columns across, five sub-columns each (``% consumed``, its STD, ``FC``, its
  STD, ``p(FDR)``), plus each drug's own ``Drug adaptive FC threshold %``. Every
  edge comes from this table.
* ``Supplementary Table 1`` — the strain dictionary: 76 organisms with a phylum
  and a culture-collection number, then the plasmids and mutants the gene work
  used, which are not organisms and stop the read.
* ``Supplementary Table 2`` — the drug dictionary: the screened compound's name,
  its parent drug's name, CAS, trade name, therapeutic indication and SMILES.
* ``Supplementary Table 13`` — the gain-of-function gene screen: 30 bacterial
  gene products x 20 of the drugs, binary.

What comes out:

* every cell the screen called a hit -> ``METABOLISES``, in
  ``taxon_drug_metabolised.csv`` — **2,575**;
* every other measured cell -> ``DOES_NOT_METABOLISE``, in
  ``taxon_drug_not_metabolised.csv`` — **17,479**. A screen measures the whole
  matrix, so a non-hit is a measurement and not an absence of curation, and it
  is its own relationship rather than a flag for the reason
  ``DOES_NOT_INHIBIT_GROWTH_OF`` and ``NO_EXCHANGE_WITH`` are;
* the four ``Control pH`` columns -> ledger rows and **no edges**. They sit
  among the strain columns and carry the same five sub-columns, and under the
  hit rule they produce 38 apparent hits between them — abiotic degradation
  read as microbial metabolism, 1,084 cells of it, if a reader took the block
  structure at face value;
* a strain no NCBI name reaches -> an ``UnresolvedTaxon`` tombstone and a ledger
  row naming the candidates that were rejected, never a guess;
* a drug no join route reaches -> a ``Drug`` node keyed on its screened name,
  ``approved = false``, ``source = zimmermann2019``. **23 of the 271**; the
  other 248 join by the screened name (195), by the file's own parent-drug name
  (43) or by the screened name with a salt suffix removed (10) — and **17 of
  those land on a node Maier 2018 minted**, which is why this prep declares
  ``maier2018`` in ``DEPENDS_ON``. A second node for a compound the other screen
  already minted would make "does this drug inhibit gut bacteria *or* get
  metabolised by them" unanswerable for it, which is the whole of D8.

**The call rule is derived, and that is the interesting part.** The sheet
publishes a per-drug depletion threshold, so half the rule is read rather than
guessed; the comparison and the significance cutoff are not published and both
change the answer. ``% consumed >= the drug's own threshold`` with
``p(FDR) <= 0.05`` reproduces the paper's own headline — **176 of 271 drugs
metabolised by at least one strain** — exactly, where ``<`` gives 175, 0.01
gives 133 and dropping the per-drug threshold for its 20% floor gives 190.
:func:`check_headline` re-derives it on every run and refuses to write when it
stops holding, because the rule decides the *type* of every edge in the source.

**Seventeen sheets are not loaded and the two that are worth naming are.**
``Supplementary Table 6`` is the metabolite matrix — 6,573 rows of
mass-to-charge *features* like ``Bisacodyl_183.0685``, with no compound
identity, no ChEBI or HMDB id and no name, so nothing in it can become a
``Metabolite`` node this graph could join to anything. ``Supplementary Table
13`` is loaded but **not as a node type**: its genes ride on the edge as
``gene_locus_tags`` and friends, because they were identified by expressing a
library in *E. coli* rather than in the 76 strains these edges come from
(docs/model.md, "no ``Gene`` node, and why").
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from microbiomekg import ontology as ont  # noqa: E402
from microbiomekg.drugs import DrugIndex, join_drug  # noqa: E402
from microbiomekg.ontology import zimmermann2019 as zm  # noqa: E402
from microbiomekg.rawdata import find_taxdump  # noqa: E402
from microbiomekg.reconcile import TaxonomyIndex, rank_depth  # noqa: E402
from microbiomekg.tables import Writer, as_list  # noqa: E402

SOURCE = zm.SOURCE

#: Reads ``drug.csv``, which ChEMBL writes and Maier appends to. Every join route
#: here reads that table, and ``maier2018`` is named as well as ``chembl``
#: because 17 of this screen's 271 compounds reach a node **only** because the
#: other screen minted one for them. Running first would mint a second node for
#: each, splitting one drug in two along exactly the seam D8 asks across.
DEPENDS_ON: list[str] = ["chembl", "maier2018"]

#: The one workbook, and the sheets read from it. Named rather than positional
#: so a re-extraction that reorders them fails loudly instead of reading the
#: wrong table.
WORKBOOK = "41586_2019_1291_MOESM1_ESM.xlsx"
SHEETS: dict[str, str] = {
    "strains": "Supplementary Table 1",
    "drugs": "Supplementary Table 2",
    "screen": "Supplementary Table 3",
    "genes": "Supplementary Table 13",
}

#: Where ``fetch.py`` puts the workbook. Under ``drug_screens/`` beside Maier
#: rather than in a directory of its own, because the two are the pair Part B's
#: W7 names and neither is the whole answer alone.
RAW_SUBDIR = Path("drug_screens") / SOURCE

#: Supplementary table 1 opens with a title line and two blanks, then a
#: sub-heading, then the 76 organisms — and then keeps going into the
#: *B. thetaiotaomicron* mutants, the cloning strains and the plasmids, which are
#: not organisms this graph has any way to key. The header is found by its first
#: cell and the organism block ends at the first cell that starts this string, so
#: neither is a row count that a re-extraction can shift.
STRAIN_HEADER_CELL = "Name"
STRAIN_SECTION_END = "Bacteroides thetaiotaomicron (Background"

#: The five sub-columns of one strain's block in the screen, in order. They are
#: read *positionally* — the label is repeated in each sub-header, so the block
#: cannot be found any other way — and one of the five decides every hit call, so
#: :func:`read_screen` checks all five rather than counting on the order.
SCREEN_COLUMN_SUBHEADS = ("% consumed", "% consumed STD", "FC", "FC STD", " p(FDR)")

EDGE_FIELDS = [
    "effect", "evidence_level", "knowledge_level", "agent_type", "primary_source",
    "source_record_id", "source_licence", "source_relation", "publications", "pmid",
    "percent_consumed", "percent_consumed_std", "fold_change", "fold_change_std",
    "fdr_p_value", "drug_threshold_percent", "incubation_hours", "replicates",
    "reported_drug_name", "parent_drug_name", "therapeutic_indication", "drug_join",
    "screen_column", "reported_name", "strain", "phylum", "reported_rank",
    "original_rank", "resolution_status", "strain_join", "taxon_join",
    "gene_locus_tags", "gene_products", "gene_protein_ids", "n_gene_products",
]

#: The same column set ``prep_maier2018.py`` writes, plus this source's three.
#: Listed in full rather than appended to whatever is already in the file so that
#: a build where Maier's raw workbooks are absent still writes a ``drug.csv``
#: whose columns match what ``blueprints/zimmermann2019.json`` declares.
DRUG_FIELDS = [
    "drug_id", "chembl_id", "pref_name", "molecule_type", "max_phase",
    "first_approval", "atc_codes", "approved", "withdrawn", "therapeutic", "oral",
    "parenteral", "topical", "smiles", "salt_form", "salt_ids", "source",
    "source_licence", "chembl_release", "prestwick_id", "pubchem_cid",
    "screen_drug_class", "screen_target_species",
    "cas", "trade_name", "therapeutic_indication",
]

#: The same shape ``unresolved_maier2018.csv`` uses. C18's accounting lives here:
#: every input cell that becomes no edge, and every column, strain, drug or gene
#: that reached something other than what it looks like it reached.
LEDGER_FIELDS = ["kind", "record_id", "subject", "detail", "reason", "source"]


def sheet_rows(path: Path, sheet: str) -> list[tuple]:
    """Every row of one named sheet, as tuples.

    Read with ``read_only=False`` for the reason ``prep_maier2018`` gives: a
    stale dimension record makes the streaming reader return a one-cell sheet,
    and that failure looks like an empty table rather than an error.
    """
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise SystemExit(
                f"{path} has no sheet named {sheet!r} "
                f"(found {', '.join(book.sheetnames)})"
            )
        return list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()


def as_float(value) -> float | None:
    """A cell as a float, or ``None`` for a blank and for anything non-numeric.

    All 20,596 measured cells of the real screen are numeric, which is worth
    knowing and is not worth trusting: a ``float(cell)`` written inline would
    raise on the first release that writes ``NA``, and a ``try`` wide enough to
    swallow that would file the pair as a measured non-hit — a measurement
    nobody made.
    """
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def read_strains(path: Path) -> list[dict]:
    """The 76 screened organisms of supplementary table 1, in sheet order.

    Each is ``{name, phylum, reference}`` — the ``Reference`` column carries a
    culture-collection number for 72 of them and the literal ``fecal isolate``
    for the other four, and it is half of the key supplementary table 3's column
    headers are matched on.
    """
    rows = sheet_rows(path, SHEETS["strains"])
    start = next(
        (i for i, row in enumerate(rows) if text(row[0]) == STRAIN_HEADER_CELL), None
    )
    if start is None:
        raise SystemExit(f"{path}: no header row whose first cell is "
                         f"{STRAIN_HEADER_CELL!r} in {SHEETS['strains']}")
    out: list[dict] = []
    for row in rows[start + 1:]:
        name = text(row[0])
        if name.startswith(STRAIN_SECTION_END):
            break
        # The sub-heading that opens the organism block has no phylum and no
        # reference; a genuine row has both columns filled.
        if not name or not text(row[1]) or not text(row[2]):
            continue
        out.append({"name": name, "phylum": text(row[1]), "reference": text(row[2])})
    return out


def read_drugs(path: Path) -> dict[str, dict]:
    """``MOLENAME -> the supplementary-table-2 row``, stopping at the blank row.

    The sheet's data is followed by a block describing each of its own columns,
    so the terminator matters: without it the dictionary gains twelve entries
    named ``MOLENAME``, ``SMILES``, ``cas`` and so on, and any count of "drugs in
    the dictionary" reads 283 against the screen's 271.
    """
    rows = sheet_rows(path, SHEETS["drugs"])
    start = next((i for i, row in enumerate(rows) if text(row[0]) == "MOLENAME"), None)
    if start is None:
        raise SystemExit(f"{path}: no MOLENAME header in {SHEETS['drugs']}")
    header = [text(c) for c in rows[start]]
    out: dict[str, dict] = {}
    for row in rows[start + 1:]:
        name = text(row[0])
        if not name:
            break
        out[name] = dict(zip(header, row))
    return out


def read_screen(path: Path) -> tuple[list[str], list[dict]]:
    """``(measured column labels in sheet order, one dict per drug row)``.

    The sheet is two header rows deep: the first names each strain once, above
    its five-column block; the second names the five. The block starts are
    therefore *the non-empty cells of the first header row*, and the five
    sub-headers are checked rather than assumed — a re-extraction that reordered
    them would otherwise put a standard deviation where a p-value goes and every
    count downstream would still look plausible.
    """
    rows = sheet_rows(path, SHEETS["screen"])
    blocks = next(
        (i for i, row in enumerate(rows) if text(row[0]) == "DrugName"), None
    )
    if blocks is None:
        raise SystemExit(f"{path}: no DrugName header row in {SHEETS['screen']}")
    header, sub = rows[blocks], rows[blocks + 1]
    starts = [(i, text(v)) for i, v in enumerate(header) if text(v)][1:]
    if not starts:
        raise SystemExit(f"{path}: {SHEETS['screen']} names no measured column")
    for at, label in starts:
        got = [text(sub[at + k]) for k in range(len(SCREEN_COLUMN_SUBHEADS))]
        want = [f"{s} {label}".strip() for s in SCREEN_COLUMN_SUBHEADS]
        if [g.replace(" ", "") for g in got] != [w.replace(" ", "") for w in want]:
            raise SystemExit(
                f"{path}: {SHEETS['screen']} column block {label!r} is "
                f"{got} where {want} was expected — the five sub-columns are "
                f"read positionally and one of them decides every hit call"
            )
    out: list[dict] = []
    for row in rows[blocks + 2:]:
        name = text(row[0])
        if not name:
            continue
        out.append({
            "drug": name,
            "threshold": as_float(row[1]),
            "cells": [
                (label, [as_float(row[at + k]) for k in range(5)])
                for at, label in starts
            ],
        })
    return [label for _at, label in starts], out


def read_genes(path: Path) -> list[dict]:
    """Supplementary table 13's gain-of-function screen, one dict per gene.

    ``{locus_tag, patric, product, protein_id, drugs}`` where ``drugs`` is the
    set of parent-drug labels the gene product metabolised. The sheet interleaves
    a *parent drug* column with one column per detected metabolite mass, and only
    the parent-drug ones are read — a mass feature is not a compound this graph
    can name (see the module docstring on supplementary table 6).
    """
    rows = sheet_rows(path, SHEETS["genes"])
    start = next(
        (i for i, row in enumerate(rows) if text(row[0]) == "Gene"), None
    )
    if start is None:
        raise SystemExit(f"{path}: no Gene header row in {SHEETS['genes']}")
    header, sub = rows[start], rows[start + 1]
    parents = [
        (i, text(v)) for i, v in enumerate(header)
        if text(v) and text(sub[i]) == "Parent drug"
    ]
    if not parents:
        raise SystemExit(f"{path}: {SHEETS['genes']} names no 'Parent drug' column")
    out: list[dict] = []
    for row in rows[start + 2:]:
        tag = text(row[0])
        if not tag:
            continue
        out.append({
            "locus_tag": tag,
            "patric": text(row[1]),
            "product": text(row[2]),
            "protein_id": text(row[3]),
            "drugs": {label for at, label in parents if row[at] == 1},
        })
    return out


def mint_id(molename: str) -> str:
    """The ``Drug`` key for a screened compound no join route reached.

    ``CAMYLOFINE DIHYDROCHLORIDE`` -> ``ZIMMERMANN2019:CAMYLOFINE-DIHYDROCHLORIDE``.
    The screen has no catalogue number to key on the way Maier's Prestwick
    library does, and the CAS column is not one either: it is multi-valued with
    inline annotations (``34381-68-5, 37517-30-9 [acebutolol]``), so parsing an
    identifier out of it would be a grammar, and a wrong parse would be a wrong
    cross-database key rather than an ugly one. The verbatim screened name is
    what the source actually asserts; it rides on the node as ``pref_name`` and
    the CAS cell rides beside it untouched.
    """
    slug = re.sub(r"[^A-Za-z0-9]+", "-", molename.strip()).strip("-").upper()
    return f"ZIMMERMANN2019:{slug}"


#: The claim :func:`check_headline` re-derives: ``(drugs, strain columns, drugs
#: metabolised by at least one strain)``, straight out of the paper's abstract.
PUBLISHED_MATRIX = (zm.SCREENED_DRUGS, zm.REPORTED_STRAINS, zm.METABOLISED_DRUGS)


def check_headline(
    screen: list[dict], columns: list[str], published: tuple[int, int, int]
) -> int:
    """Re-derive the call rule against a published headline, or refuse to write.

    The sheet publishes each drug's depletion threshold but no significance
    cutoff and no comparison, and :func:`.zimmermann2019.relation_for` decides
    the *type* of every edge in this source — so the rule is checked on every run
    rather than trusted from the module. A mismatch is a hard stop: writing the
    tables anyway would invert an unknown share of 20,054 edges silently, which
    is the failure this loader is least able to notice afterwards.

    **The shape is checked before the count, and both are hard stops.** A gate
    that quietly stopped applying because a re-extraction dropped a column would
    be worse than no gate — and that is the whole reason ``published`` is an
    argument rather than a module constant read in here. ``--published-matrix``
    defaults to :data:`PUBLISHED_MATRIX`; the test fixtures pass their own
    triple, so the reduced workbook runs *this* code path with *this* arithmetic
    instead of a version of the check that turns itself off.

    Returns the number of drugs metabolised by at least one strain.
    """
    drugs, strain_columns, headline = published
    strains = [label for label in columns if not zm.is_control_column(label)]
    if (len(screen), len(strains)) != (drugs, strain_columns):
        raise SystemExit(
            f"the screen is {len(screen)} drugs x {len(strains)} strain columns "
            f"(plus {len(columns) - len(strains)} controls), not "
            f"{drugs} x {strain_columns}. Every published number this loader "
            f"checks itself against is for that matrix."
        )
    metabolised = sum(
        1 for row in screen
        if any(
            zm.relation_for(values[0], row["threshold"], values[4])
            == zm.RELATION_METABOLISES
            for label, values in row["cells"]
            if not zm.is_control_column(label)
        )
    )
    if metabolised != headline:
        raise SystemExit(
            f"the call rule makes {metabolised} of {len(screen)} drugs "
            f"metabolised by at least one strain; {headline} was published. "
            f"The rule is derived, not stated: re-derive it before writing, "
            f"because it decides which relationship every cell becomes."
        )
    return metabolised


def gene_column_for(
    gene: dict, idx: TaxonomyIndex, columns: dict[str, dict], rank_ceiling: str
) -> tuple[str, str]:
    """``(the screened column this gene's organism is, or "", why not)``.

    The PATRIC feature id carries the donor genome's **NCBI taxid**, which is the
    only identifier the gene sheet and the screen share. Two things have to hold
    before a gene's evidence may be attached to an edge, and neither is optional:

    * the taxid must promote to the same species as the screened column, and
    * the column's own culture-collection number must appear inside that strain
      taxon's name.

    The second is what stops a *B. thetaiotaomicron* gene from being attributed
    to all three screened *B. thetaiotaomicron* isolates, and it is why the match
    is not made on the strain's *name*: NCBI has since renamed 483217 to
    ``Phocaeicola dorei DSM 17855`` while the sheet still says *Bacteroides*, so
    the binomial no longer agrees and the collection number still does.

    Anything other than exactly one match is a ledger row rather than a guess.
    """
    tax_id = zm.patric_tax_id(gene["patric"])
    if tax_id is None:
        return "", f"{gene['patric']!r} is not a PATRIC feature id"
    # The same ceiling the screened columns were resolved at, not this source's
    # `RANK_CEILING` guard: the two ids have to be comparable, and a gene
    # promoted to the genus would match every screened column of that genus.
    res = idx.resolve(tax_id=tax_id, rank_ceiling=rank_ceiling)
    if res.tax_id is None:
        return "", f"PATRIC genome taxid {tax_id} reaches no taxon: {res.note}"
    strain_key = zm.join_key(idx.scientific_name.get(tax_id, ""))
    hits = [
        label for label, column in columns.items()
        if column["tax_id"] == res.tax_id
        and column["reference_key"]
        and column["reference_key"] in strain_key
    ]
    if len(hits) != 1:
        return "", (
            f"PATRIC genome taxid {tax_id} matches {len(hits)} screened columns "
            f"of taxon {res.tax_id}"
        )
    return hits[0], ""


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--raw", type=Path, default=Path("data/raw"))
    ap.add_argument("--tables", type=Path, default=None,
                    help=f"directory holding {WORKBOOK} "
                         f"(default: <raw>/{RAW_SUBDIR}).")
    ap.add_argument("--taxdump", type=Path, default=None)
    ap.add_argument("--out", type=Path, default=Path("data/csv"))
    ap.add_argument("--rank-ceiling", default="species",
                    help="Rank strains and subspecies are promoted to before keying.")
    ap.add_argument("--published-matrix", type=_triple, default=PUBLISHED_MATRIX,
                    help="DRUGS,STRAINS,METABOLISED — the published claim this "
                         "run must reproduce before it writes anything. Defaults "
                         "to the paper's. Only the test fixtures pass another, "
                         "and they pass one rather than switching the check off.")
    args = ap.parse_args(argv)

    tables = args.tables or (args.raw / RAW_SUBDIR)
    workbook = tables / WORKBOOK
    if not workbook.is_file():
        # Exit 3, not 2 — "this source's raw file is not on this machine" rather
        # than "this script was called wrong", so an absent workbook leaves the
        # build without failing it.
        print(f"no Zimmermann 2019 workbook at {workbook}", file=sys.stderr)
        return 3
    try:
        taxdump = args.taxdump or find_taxdump(args.raw)
    except FileNotFoundError as e:
        ap.error(str(e))

    out = args.out
    strains = read_strains(workbook)
    library = read_drugs(workbook)
    columns, screen = read_screen(workbook)
    genes = read_genes(workbook)
    controls = [label for label in columns if zm.is_control_column(label)]
    print(f"read {len(screen):,} drugs x {len(columns)} measured columns = "
          f"{len(screen) * len(columns):,} cells from {workbook.name}")
    print(f"  {len(strains):,} strains and {len(library):,} drugs in the "
          f"dictionaries, {len(genes):,} gene products in the gain-of-function "
          f"screen")
    print(f"  {len(controls)} of the measured columns are abiotic controls, not "
          f"organisms: {', '.join(label.strip() for label in controls)}")
    metabolised = check_headline(screen, columns, args.published_matrix)
    print(f"  the call rule reproduces the published headline: {metabolised} of "
          f"{len(screen)} drugs ({100 * metabolised / len(screen):.0f}%) "
          f"metabolised by at least one strain")

    index = DrugIndex.from_csv(out / "drug.csv", exclude_source=SOURCE)
    print(f"drug.csv: {len(index.names):,} names this screen may join to")

    print(f"loading taxdump from {taxdump} ...", flush=True)
    idx = TaxonomyIndex.from_taxdump(taxdump)
    print(f"  {len(idx.parent):,} taxa, {len(idx.names):,} name keys", flush=True)

    edges = {
        zm.RELATION_METABOLISES: Writer(
            out / "taxon_drug_metabolised.csv", ["tax_id", "drug_id", *EDGE_FIELDS],
            dedupe_full=True, merge=True, owner=("primary_source", SOURCE),
        ),
        zm.RELATION_NO_METABOLISM: Writer(
            out / "taxon_drug_not_metabolised.csv",
            ["tax_id", "drug_id", *EDGE_FIELDS],
            dedupe_full=True, merge=True, owner=("primary_source", SOURCE),
        ),
    }
    drugs = Writer(
        out / "drug.csv", DRUG_FIELDS,
        key="drug_id", merge=True, owner=("source", SOURCE),
    )
    unresolved_nodes = Writer(
        out / "unresolved_taxa.csv",
        ["unresolved_id", "raw_name", "reported_rank", "original_rank",
         "reported_tax_id", "source", "status", "candidates", "note", "n_signatures"],
        key="unresolved_id", merge=True, owner=("source", SOURCE),
    )
    ledger = Writer(
        out / f"unresolved_{SOURCE}.csv", LEDGER_FIELDS,
        merge=True, owner=("source", SOURCE),
    )
    cited = Writer(
        out / "cited_taxa.csv", ["tax_id", "source", "n_signatures"],
        key=("tax_id", "source"), merge=True, owner=("source", SOURCE),
    )

    def note(kind: str, record_id: str, subject: str, detail: str, reason: str) -> None:
        ledger.add({"kind": kind, "record_id": record_id, "subject": subject,
                    "detail": detail, "reason": reason, "source": SOURCE})

    # ------------------------------------------------- the columns and strains
    #: normalised ``Name`` + ``Reference`` -> the table-1 row, plus the ``Name``
    #: alone for the organisms whose name is unique in the sheet. The name-only
    #: key is what reaches the four fecal isolates, whose column header carries
    #: no designation at all; it is *not* added for a repeated name, because
    #: seven rows say ``Bacteroides fragilis`` and a header matching all of them
    #: would attach one isolate's row to another's strain designation.
    seen_names: Counter[str] = Counter(zm.join_key(s["name"]) for s in strains)
    by_collection: dict[str, dict] = {}
    by_name: dict[str, dict] = {}
    for strain in strains:
        by_collection.setdefault(
            zm.join_key(strain["name"] + strain["reference"]), strain
        )
        if seen_names[zm.join_key(strain["name"])] == 1:
            by_name.setdefault(zm.join_key(strain["name"]), strain)

    resolved: dict[str, dict] = {}
    resolutions: Counter[str] = Counter()
    strain_joins: Counter[str] = Counter()
    taxon_joins: Counter[str] = Counter()
    unresolved_hits: Counter[str] = Counter()
    for label in columns:
        if zm.is_control_column(label):
            note("column", label.strip(), label.strip(), "",
                 "an abiotic degradation control, not an organism: loading it "
                 "would write chemistry as microbial metabolism")
            continue
        override = zm.COLUMN_OVERRIDES.get(label)
        key = zm.join_key(override or label)
        entry, route = by_collection.get(key), "collection"
        if entry is None:
            entry, route = by_name.get(key), "name"
        if override:
            route = "override"
        if entry is None:
            note("column", label.strip(), label.strip(), "",
                 "no supplementary-table-1 row spells this screened column")
            continue
        strain_joins[route] += 1
        reported = entry["name"]
        species = zm.SPECIES_OVERRIDES.get(reported.casefold())
        taxon_route = "override" if species else "exact"
        res = idx.resolve(species or reported, rank_ceiling=args.rank_ceiling)
        resolutions[res.status] += 1
        if res.tax_id is None:
            # The taxa this name *could* have meant, where the decision was made
            # deliberately rather than by failing: a tombstone that says only
            # "no name match" is a dead end, and one that names what was rejected
            # is reversible.
            rejected = zm.AMBIGUOUS_STRAINS.get(reported.casefold(), ())
            uid = f"unresolved:{SOURCE}:{reported.casefold()}"
            unresolved_nodes.add({
                "unresolved_id": uid, "raw_name": reported,
                "reported_rank": zm.REPORTED_RANK,
                "original_rank": res.original_rank or "", "reported_tax_id": "",
                "source": SOURCE, "status": res.status,
                "candidates": as_list(
                    str(c) for c in (res.candidates or [t for t, _n in rejected])
                ),
                "note": res.note, "n_signatures": "0",
            })
            unresolved_hits[uid] += 1
            note("strain", label.strip(), reported,
                 "; ".join(f"{name} ({tid})" for tid, name in rejected)
                 or entry["reference"],
                 f"taxon {res.status}: {res.note}"
                 + (" — NCBI holds more than one candidate and nothing in the "
                    "row chooses between them, so no override is written"
                    if rejected else ""))
            continue
        rank = idx.rank.get(res.tax_id) or ""
        own, ceiling = rank_depth(rank), rank_depth(zm.RANK_CEILING)
        if own is None or ceiling is None or own < ceiling:
            # Every organism here is one cultured isolate, so a resolution this
            # broad is a surprise rather than a filter — and the ledger row
            # carries the id and rank it did reach, so the ceiling is reversible.
            note("strain", label.strip(), reported,
                 f"{res.tax_id} ({rank or 'unplaced'})",
                 f"resolved rank is broader than {zm.RANK_CEILING!r}")
            continue
        taxon_joins[taxon_route] += 1
        resolved[label] = {
            "tax_id": res.tax_id, "reported_name": reported,
            "strain": entry["reference"], "phylum": entry["phylum"],
            "reference_key": zm.join_key(entry["reference"]),
            "original_rank": res.original_rank or rank,
            "resolution_status": res.status,
            "strain_join": route, "taxon_join": taxon_route,
        }

    # ------------------------------------------------------------- the drugs
    drug_joins: Counter[str] = Counter()
    joined_via: Counter[str] = Counter()
    minted: dict[str, str] = {}
    drug_ids: dict[str, str] = {}
    drug_routes: dict[str, str] = {}
    for row in screen:
        molename = row["drug"]
        entry = library.get(molename)
        if entry is None:
            note("drug", molename, molename, "",
                 "screened drug with no supplementary-table-2 row: it carries no "
                 "parent name, CAS or indication, and only the verbatim route "
                 "can reach a node")
            entry = {}
        parent = text(entry.get("name"))
        drug_id, route, others = join_drug(zm.drug_variants(molename, parent), index)
        if others:
            # The screened name reaches a ChEMBL *salt* node while the parent or
            # salt-stripped spelling reaches the parent molecule. That is
            # ChEMBL's documented parent gap showing through, not a defect here —
            # but a precedence rule that hid it would leave the graph's own count
            # unreadable.
            note("drug", molename, molename,
                 f"{route} -> {drug_id}; also {','.join(others)}",
                 "two join routes reached different Drug nodes; the route the "
                 "screen's own spelling took wins")
        if not drug_id:
            drug_id = mint_id(molename)
            if drug_id not in minted:
                minted[drug_id] = molename
                drugs.add({
                    "drug_id": drug_id, "chembl_id": "", "pref_name": molename,
                    "molecule_type": "", "max_phase": "", "first_approval": "",
                    "atc_codes": "",
                    # Not `true`: every compound here is orally administered and
                    # marketed, but `approved` in this graph means ChEMBL
                    # max_phase 4, and inheriting the claim from a screening
                    # library would put an unverified regulatory status on 23
                    # nodes.
                    "approved": "false", "withdrawn": "false", "therapeutic": "",
                    "oral": "", "parenteral": "", "topical": "",
                    "smiles": text(entry.get("SMILES")),
                    "salt_form": "false", "salt_ids": "", "source": SOURCE,
                    "source_licence": ont.SOURCE_LICENCE.get(SOURCE, ""),
                    "chembl_release": "", "prestwick_id": "", "pubchem_cid": "",
                    "screen_drug_class": "", "screen_target_species": "",
                    "cas": text(entry.get("cas")),
                    "trade_name": text(entry.get("TradeName")),
                    "therapeutic_indication": text(entry.get("TherapeuticIndication")),
                })
        else:
            joined_via[index.source_of.get(drug_id, "")] += 1
        drug_joins[route] += 1
        drug_ids[molename] = drug_id
        drug_routes[molename] = route

    # -------------------------------------------------------------- the genes
    #: ``(screened column, normalised drug name) -> the gene products``.
    gene_evidence: dict[tuple[str, str], list[dict]] = {}
    unplaced_genes = 0
    for gene in genes:
        label, why = gene_column_for(gene, idx, resolved, args.rank_ceiling)
        if not label:
            unplaced_genes += 1
            note("gene", gene["locus_tag"], gene["product"], gene["patric"], why)
            continue
        for drug in sorted(gene["drugs"]):
            gene_evidence.setdefault((label, zm.join_key(drug)), []).append(gene)

    # -------------------------------------------------------------- the cells
    counters: Counter[str] = Counter()
    per_relation: Counter[str] = Counter()
    taxa_seen: "OrderedDict[int, int]" = OrderedDict()
    gene_pairs_used: set[tuple[str, str]] = set()
    gene_on_relation: Counter[str] = Counter()
    for row in screen:
        molename = row["drug"]
        drug_id = drug_ids.get(molename)
        entry = library.get(molename, {})
        drug_key = zm.join_key(molename)
        for label, values in row["cells"]:
            counters["cells"] += 1
            if zm.is_control_column(label):
                counters["control_column"] += 1
                continue
            strain = resolved.get(label)
            if strain is None:
                counters["unusable_strain"] += 1
                continue
            relationship = zm.relation_for(values[0], row["threshold"], values[4])
            if relationship is None:
                counters["not_measured"] += 1
                note("cell", f"{molename}|{label.strip()}", molename,
                     strain["reported_name"],
                     "the screen leaves a number out for this pair: measured as "
                     "neither a hit nor a non-hit")
                continue
            if drug_id is None:
                counters["unusable_drug"] += 1
                continue
            effect, source_relation = zm.SOURCE_RELATIONS[relationship]
            found = gene_evidence.get((label, drug_key), [])
            if found:
                gene_pairs_used.add((label, drug_key))
                gene_on_relation[relationship] += 1
            tax_id = strain["tax_id"]
            taxa_seen[tax_id] = taxa_seen.get(tax_id, 0) + 1
            per_relation[relationship] += 1
            counters["edges"] += 1
            edges[relationship].add({
                "tax_id": str(tax_id),
                "drug_id": drug_id,
                "effect": effect,
                "evidence_level": zm.EVIDENCE_LEVEL,
                "knowledge_level": ont.knowledge_level(SOURCE),
                "agent_type": ont.agent_type(SOURCE),
                "primary_source": SOURCE,
                "source_record_id": f"{SOURCE}:{molename}|{label.strip()}",
                "source_licence": ont.SOURCE_LICENCE.get(SOURCE, ""),
                "source_relation": source_relation,
                "publications": as_list([zm.PUBLICATION]),
                "pmid": str(zm.PUBMED_ID),
                "percent_consumed": _num(values[0]),
                "percent_consumed_std": _num(values[1]),
                "fold_change": _num(values[2]),
                "fold_change_std": _num(values[3]),
                "fdr_p_value": _num(values[4]),
                "drug_threshold_percent": _num(row["threshold"]),
                "incubation_hours": repr(zm.INCUBATION_HOURS),
                "replicates": str(zm.REPLICATES),
                "reported_drug_name": molename,
                "parent_drug_name": text(entry.get("name")),
                "therapeutic_indication": text(entry.get("TherapeuticIndication")),
                "drug_join": drug_routes.get(molename, ""),
                "screen_column": label.strip(),
                "reported_name": strain["reported_name"],
                "strain": strain["strain"],
                "phylum": strain["phylum"],
                "reported_rank": zm.REPORTED_RANK,
                "original_rank": strain["original_rank"],
                "resolution_status": strain["resolution_status"],
                "strain_join": strain["strain_join"],
                "taxon_join": strain["taxon_join"],
                "gene_locus_tags": as_list(g["locus_tag"] for g in found),
                "gene_products": as_list(g["product"] for g in found),
                "gene_protein_ids": as_list(g["protein_id"] for g in found),
                "n_gene_products": str(len(found)) if found else "",
            })

    for pair, found in sorted(gene_evidence.items()):
        if pair not in gene_pairs_used:
            note("gene", f"{found[0]['patric']}|{pair[1]}", pair[0],
                 "|".join(g["locus_tag"] for g in found),
                 "the gain-of-function screen names a gene for this pair but the "
                 "76-strain screen measured no cell for it")

    for row in unresolved_nodes.rows:
        if row["source"] == SOURCE:
            row["n_signatures"] = str(unresolved_hits.get(row["unresolved_id"], 0))
    for tid, n in sorted(taxa_seen.items()):
        cited.add({"tax_id": str(tid), "source": SOURCE, "n_signatures": str(n)})

    tables_out = (*edges.values(), drugs, unresolved_nodes, ledger, cited)
    counts = {w.path.name: w.flush() for w in tables_out}

    print(f"\nread {counters['cells']:,} cells -> {counters['edges']:,} edges over "
          f"{len(taxa_seen):,} taxa and {len(drug_ids):,} drugs")
    print("  by relationship: " + ", ".join(
        f"{rel} {per_relation[rel]:,}" for rel in zm.METABOLISM_RELATIONSHIPS))
    print(f"  measured non-hits kept as {zm.RELATION_NO_METABOLISM}: "
          f"{per_relation[zm.RELATION_NO_METABOLISM]:,} (never folded into the "
          f"hit edge)")
    print(f"  pairs carrying a gain-of-function gene product: "
          f"{sum(gene_on_relation.values()):,} (" + ", ".join(
              f"{rel} {gene_on_relation[rel]:,}"
              for rel in zm.METABOLISM_RELATIONSHIPS) + ")")
    print("  drug join: " + ", ".join(f"{r} {n:,}" for r, n in drug_joins.most_common()))
    print("  joined to a node written by: " + ", ".join(
        f"{s or 'unknown'} {n:,}" for s, n in joined_via.most_common()))
    print("  strain join: " + ", ".join(
        f"{r} {n:,}" for r, n in strain_joins.most_common()))
    print("  organism resolution: " + ", ".join(
        f"{s} {n:,}" for s, n in resolutions.most_common()))
    print(f"  not loaded: {counters['not_measured']:,} cells the screen left a "
          f"number out of, {counters['control_column']:,} on an abiotic control "
          f"column, {counters['unusable_strain']:,} on a strain that reached no "
          f"taxon, {counters['unusable_drug']:,} on a drug with no identifier")
    print(f"  minted {len(minted):,} Drug nodes for screened compounds no join "
          f"route reached")
    print(f"  gene products the screen's own columns could not place: "
          f"{unplaced_genes:,} of {len(genes):,}")
    for name_, n in counts.items():
        print(f"  {name_:34s} {n:>9,}")
    shared = {w.path.name: w.merged_in for w in tables_out if w.merged_in}
    if shared:
        print("  merged into tables another source had written: " + ", ".join(
            f"{name_} +{n:,}" for name_, n in shared.items()))
    return 0


def _triple(value: str) -> tuple[int, int, int]:
    """``"271,76,176"`` -> ``(271, 76, 176)``."""
    parts = [p.strip() for p in str(value).split(",")]
    if len(parts) != 3 or not all(p.isdigit() for p in parts):
        raise argparse.ArgumentTypeError(
            f"{value!r} is not DRUGS,STRAINS,METABOLISED"
        )
    a, b, c = (int(p) for p in parts)
    return a, b, c


def _num(value: float | None) -> str:
    """A measurement as a bare number string, or ``""``.

    The blueprint declares these columns ``float``; a stray unit or comparison
    operator would pin the whole property to ``string`` for the graph's lifetime
    (C16), and an empty cell must stay empty rather than become a zero that reads
    as "measured, no depletion".
    """
    return "" if value is None else repr(value)


if __name__ == "__main__":
    raise SystemExit(main())
